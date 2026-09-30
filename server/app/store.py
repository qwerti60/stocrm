from __future__ import annotations

import json
import threading
import time
import uuid
from pathlib import Path
from typing import Any

from app.config import settings

_LOCK = threading.Lock()
_PATH = Path(__file__).resolve().parent.parent / "data" / "week3.json"

_EMPTY: dict[str, Any] = {
    "chats": {},
    "vin": [],
    "bookings": [],
    "inbox": {},
    "sessions": {},
    "bonuses": {},
    "warned": [],
    "campaigns": [],
    "shorts": [],
    "hidden_shorts": [],
    "audience": {},
    "identities": {},
    "hidden_cars": {},
    "garage": {},
    "devices": {},
}

MEDIA = Path(__file__).resolve().parent.parent / "data" / "media"

DEFAULT_PROMOS = [
    {
        "id": "seed-oil",
        "title": "Замена масла",
        "subtitle": "от 4 990 ₽ · масло + фильтр + работа",
        "badge": "Акция",
        "pinned": True,
        "kind": "promo",
    },
    {
        "id": "seed-diag",
        "title": "Диагностика VAG бесплатно",
        "subtitle": "При первом визите через приложение",
        "badge": "В приложении",
        "pinned": True,
        "kind": "promo",
    },
    {
        "id": "seed-tires",
        "title": "−20% на шиномонтаж",
        "subtitle": "До конца месяца, любой филиал",
        "badge": "Акция",
        "pinned": True,
        "kind": "promo",
    },
]

DEFAULT_SHORTS = [
    {"id": "s1", "title": "Ваш VAG в надёжных руках", "subtitle": "Сервис · запчасти · забота о VAG", "url": "", "poster": ""},
    {"id": "s2", "title": "Машина готова!", "subtitle": "Push, когда статус в STOCRM сменится", "url": "", "poster": ""},
    {"id": "s3", "title": "Замена масла от 4 990 ₽", "subtitle": "Масло + фильтр + работа", "url": "", "poster": ""},
    {"id": "s4", "title": "Подбор запчастей по VIN", "subtitle": "Заявка уходит менеджеру в админку", "url": "", "poster": ""},
]


def _load() -> dict[str, Any]:
    if not _PATH.is_file():
        return json.loads(json.dumps(_EMPTY))
    try:
        data = json.loads(_PATH.read_text(encoding="utf-8"))
    except Exception:
        return json.loads(json.dumps(_EMPTY))
    for k, v in _EMPTY.items():
        data.setdefault(k, json.loads(json.dumps(v)))
    return data


def _save(data: dict[str, Any]) -> None:
    _PATH.parent.mkdir(parents=True, exist_ok=True)
    tmp = _PATH.with_suffix(".tmp")
    tmp.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(_PATH)


def thread_key(contact_id: Any, phone: str | None = None) -> str:
    try:
        cid = int(contact_id or 0)
    except (TypeError, ValueError):
        cid = 0
    if cid:
        return str(cid)
    return f"phone:{phone or '?'}"


def list_messages(key: str) -> list[dict[str, Any]]:
    with _LOCK:
        data = _load()
        return list(data["chats"].get(key, []))


def add_message(
    key: str,
    *,
    text: str,
    from_staff: bool,
    staff_name: str = "Егор",
    kind: str = "chat",
    phone: str = "",
    name: str = "",
) -> dict[str, Any]:
    msg = {
        "id": uuid.uuid4().hex[:12],
        "text": text.strip(),
        "from_staff": from_staff,
        "staff_name": staff_name if from_staff else "",
        "kind": kind,
        "created": int(time.time()),
        "phone": phone,
        "name": name,
    }
    with _LOCK:
        data = _load()
        data["chats"].setdefault(key, [])
        data["chats"][key].append(msg)
        meta = data.setdefault("meta", {})
        meta[key] = {"phone": phone or meta.get(key, {}).get("phone", ""), "name": name or meta.get(key, {}).get("name", "")}
        _save(data)
    return msg


def list_threads() -> list[dict[str, Any]]:
    with _LOCK:
        data = _load()
        meta = data.get("meta") or {}
        out = []
        for key, msgs in data["chats"].items():
            last = msgs[-1] if msgs else {}
            info = meta.get(key) or {}
            out.append(
                {
                    "key": key,
                    "phone": info.get("phone") or last.get("phone") or "",
                    "name": info.get("name") or last.get("name") or "",
                    "count": len(msgs),
                    "last": last.get("text") or "",
                    "last_at": last.get("created") or 0,
                    "from_staff": last.get("from_staff"),
                }
            )
        out.sort(key=lambda x: x["last_at"], reverse=True)
        return out


def add_vin(key: str, *, vin: str, part: str, phone: str = "", name: str = "") -> dict[str, Any]:
    row = {
        "id": uuid.uuid4().hex[:10],
        "kind": "vin",
        "key": key,
        "vin": vin.strip().upper(),
        "part": part.strip(),
        "phone": phone,
        "name": name,
        "created": int(time.time()),
        "status": "new",
    }
    with _LOCK:
        data = _load()
        data["vin"].insert(0, row)
        _save(data)
    return row


def list_vin(key: str | None = None) -> list[dict[str, Any]]:
    with _LOCK:
        rows = list(_load()["vin"])
    if key:
        return [r for r in rows if str(r.get("key")) == str(key)]
    return rows


def add_booking_ticket(
    key: str,
    *,
    branch: str = "",
    branch_id: str = "",
    when: str = "",
    comment: str = "",
    car: str = "",
    phone: str = "",
    name: str = "",
    offer_id: Any = None,
) -> dict[str, Any]:
    row = {
        "id": uuid.uuid4().hex[:10],
        "kind": "book",
        "key": key,
        "branch": branch,
        "branch_id": str(branch_id or ""),
        "when": when,
        "comment": comment,
        "car": car,
        "phone": phone,
        "name": name,
        "offer_id": offer_id,
        "created": int(time.time()),
        "status": "new",
    }
    with _LOCK:
        data = _load()
        data.setdefault("bookings", []).insert(0, row)
        _save(data)
    return row


def list_bookings(key: str | None = None) -> list[dict[str, Any]]:
    with _LOCK:
        rows = list(_load().get("bookings") or [])
    if key:
        return [r for r in rows if str(r.get("key")) == str(key)]
    return rows


def list_tickets(key: str | None = None) -> list[dict[str, Any]]:
    rows = list_vin(key) + list_bookings(key)
    rows.sort(key=lambda x: int(x.get("created") or 0), reverse=True)
    return rows


def put_session(token: str, user: dict[str, Any]) -> None:
    if not token:
        return
    with _LOCK:
        data = _load()
        data.setdefault("sessions", {})[token] = dict(user)
        _save(data)


def get_session(token: str) -> dict[str, Any] | None:
    if not token:
        return None
    with _LOCK:
        row = (_load().get("sessions") or {}).get(token)
    return dict(row) if isinstance(row, dict) else None


def drop_session(token: str) -> None:
    with _LOCK:
        data = _load()
        (data.get("sessions") or {}).pop(token, None)
        _save(data)


def all_sessions() -> dict[str, dict[str, Any]]:
    with _LOCK:
        raw = _load().get("sessions") or {}
    return {str(k): dict(v) for k, v in raw.items() if isinstance(v, dict)}


def save_inbox_note(contact_id: int, note: dict[str, Any]) -> None:
    key = str(int(contact_id))
    with _LOCK:
        data = _load()
        box = data.setdefault("inbox", {}).setdefault(key, [])
        box.insert(0, note)
        data["inbox"][key] = box[:200]
        _save(data)


def load_inbox(contact_id: int) -> list[dict[str, Any]]:
    key = str(int(contact_id))
    with _LOCK:
        return list((_load().get("inbox") or {}).get(key) or [])


def mark_inbox_read(contact_id: int, note_id: str) -> None:
    key = str(int(contact_id))
    with _LOCK:
        data = _load()
        for n in data.setdefault("inbox", {}).setdefault(key, []):
            if n.get("id") == note_id:
                n["read"] = True
                break
        _save(data)


def parse_phones(*chunks: str) -> list[str]:
    import re

    found: list[str] = []
    seen: set[str] = set()
    blob = "\n".join(c or "" for c in chunks)
    for m in re.findall(r"\d{10,15}", blob):
        d = re.sub(r"\D", "", m)
        if d.startswith("8") and len(d) == 11:
            d = "7" + d[1:]
        if len(d) == 10:
            d = "7" + d
        if d.startswith("7") and len(d) == 11 and d not in seen:
            seen.add(d)
            found.append(d)
    return found


def phones_from_table(raw: bytes, filename: str = "") -> list[str]:
    name = (filename or "").lower()
    text = ""
    if name.endswith(".xlsx") or raw[:2] == b"PK":
        try:
            import io

            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            parts: list[str] = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    for cell in row:
                        if cell is not None:
                            parts.append(str(cell))
            text = "\n".join(parts)
        except Exception:
            text = raw.decode("utf-8", errors="ignore")
    else:
        text = raw.decode("utf-8-sig", errors="ignore")
    return parse_phones(text)


def audience_by_phone(phone: str) -> dict[str, Any] | None:
    want = (phone or "").strip()
    if not want:
        return None
    for row in list_audience():
        if str(row.get("phone") or "") == want:
            return row
    ident = get_identity(want)
    if ident:
        return ident
    return None


def save_promo_image(cid: str, data: bytes, suffix: str = ".jpg") -> str:
    MEDIA.mkdir(parents=True, exist_ok=True)
    ext = suffix if suffix.startswith(".") else f".{suffix}"
    if ext.lower() not in (".jpg", ".jpeg", ".png", ".webp"):
        ext = ".jpg"
    path = MEDIA / f"{cid}{ext}"
    path.write_bytes(data)
    return str(path.name)


def promo_image_path(cid: str) -> Path | None:
    if not cid:
        return None
    for p in MEDIA.glob(f"{cid}.*"):
        if p.is_file():
            return p
    return None


def money(v: Any) -> int:
    if v is None or v is False:
        return 0
    s = str(v).replace(" ", "").replace("\xa0", "").replace(",", ".")
    try:
        return int(round(float(s)))
    except (TypeError, ValueError):
        return 0


def bonus_points(amount: int, board_id: int = 0) -> int:
    """5% как в кэшбэке CRM. Мелкие ЗН чужих воронок (Эрвье и т.п.) не копятся."""
    if amount <= 0:
        return 0
    pct = max(int(settings.bonus_percent or 5), 1)
    min_sum = max(int(settings.bonus_min_sum or 0), 0)
    main = int(settings.stocrm_board_id or 1097)
    try:
        board = int(board_id or 0)
    except (TypeError, ValueError):
        board = 0
    if min_sum and amount < min_sum and board and board != main:
        return 0
    return int(round(amount * pct / 100))


def is_bonus_offer(status_id: int, title: str) -> bool:
    """Как в CRM: только «Успешно реализовано». «Выполнен» — ещё не финал воронки."""
    t = (title or "").lower()
    if status_id in (4, 5, 5648) or any(x in t for x in ("отказ", "мусор", "не приехал")):
        return False
    return status_id == 6 or "успешн" in t


def _offer_ts(o: dict[str, Any]) -> int:
    for key in (
        "OFFER_FIRST_FINAL_STATUS_FRONTEND_TIMESTAMP",
        "OFFER_FIRST_COMPLETE_STATUS_DATE_FRONTEND_TIMESTAMP",
        "OFFER_STATUS_UPDATE_FRONTEND_TIMESTAMP",
        "OFFER_DATE_UPDATE_FRONTEND_TIMESTAMP",
        "OFFER_DATE_CREATE_FRONTEND_TIMESTAMP",
    ):
        v = o.get(key)
        if v in (None, False, "", 0, "0"):
            continue
        try:
            ts = int(v)
        except (TypeError, ValueError):
            continue
        if ts > 0:
            return ts
    return int(time.time())


def sync_bonuses(contact_id: int, offers: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Пересобрать 5% с успешно реализованных ЗН CRM. Возвращает новые начисления."""
    pct = max(int(settings.bonus_percent or 5), 1)
    expire_days = max(int(settings.bonus_expire_days or 365), 1)
    key = str(int(contact_id))
    created: list[dict[str, Any]] = []
    with _LOCK:
        data = _load()
        book: dict[str, Any] = data["bonuses"].setdefault(key, {"events": []})
        spent = [e for e in book.get("events") or [] if int(e.get("delta") or 0) < 0]
        old_acc_ids = {
            str(e.get("id"))
            for e in book.get("events") or []
            if int(e.get("delta") or 0) > 0
        }
        seen: set[str] = set()
        accrued: list[dict[str, Any]] = []
        for o in offers:
            oid = o.get("OFFER_ID") or o.get("id")
            if oid is None:
                continue
            try:
                sid = int(o.get("OFFER_STATUS_ID") or o.get("status_id") or 0)
            except (TypeError, ValueError):
                sid = 0
            title = str(o.get("STATUS_NAME") or o.get("status") or "").strip()
            if not is_bonus_offer(sid, title):
                continue
            amount = money(
                o.get("OFFER_SUM") or o.get("WORKS_SUM") or o.get("ORDERS_SUM") or o.get("sum") or o.get("works_sum")
            )
            if amount <= 0:
                continue
            if str(oid) in seen:
                continue
            seen.add(str(oid))
            try:
                board_id = int(o.get("BOARD_ID") or o.get("board_id") or 0)
            except (TypeError, ValueError):
                board_id = 0
            delta = bonus_points(amount, board_id)
            if delta <= 0:
                continue
            when = _offer_ts(o)
            ev = {
                "id": f"zn-{oid}",
                "title": f"Начисление {pct}%",
                "delta": delta,
                "when": when,
                "note": f"ЗН № {oid} · {title} · {amount} ₽",
                "offer_id": oid,
                "expires": when + expire_days * 86400,
                "source": "crm",
            }
            accrued.append(ev)
        accrued.sort(key=lambda e: int(e.get("when") or 0), reverse=True)
        book["events"] = accrued + spent
        created = [e for e in accrued if str(e.get("id")) not in old_acc_ids]
        _save(data)
    return created


def _bonus_expires(e: dict[str, Any], expire_days: int) -> int:
    try:
        exp = int(e.get("expires") or 0)
    except (TypeError, ValueError):
        exp = 0
    if exp > 0:
        return exp
    try:
        when = int(e.get("when") or 0)
    except (TypeError, ValueError):
        when = 0
    if when > 0:
        return when + expire_days * 86400
    return 0


def _live_bonus(events: list[dict[str, Any]], now: int | None = None) -> dict[str, Any]:
    """Текущий остаток: FIFO-списания, затем сгорание по сроку."""
    expire_days = max(int(settings.bonus_expire_days or 365), 1)
    warn_days = max(int(settings.bonus_warn_days or 5), 1)
    now = int(now or time.time())
    lots: list[dict[str, Any]] = []
    ordered = sorted(
        enumerate(events),
        key=lambda item: (int(item[1].get("when") or 0), item[0]),
    )
    for _, e in ordered:
        try:
            delta = int(e.get("delta") or 0)
        except (TypeError, ValueError):
            delta = 0
        if delta > 0:
            lots.append({"event": e, "remaining": delta, "expires": _bonus_expires(e, expire_days)})
            continue
        if delta >= 0:
            continue
        need = -delta
        try:
            spend_at = int(e.get("when") or now)
        except (TypeError, ValueError):
            spend_at = now
        for lot in lots:
            if need <= 0:
                break
            exp = int(lot["expires"] or 0)
            if exp and exp <= spend_at:
                continue
            if int(lot["remaining"]) <= 0:
                continue
            take = min(int(lot["remaining"]), need)
            lot["remaining"] = int(lot["remaining"]) - take
            need -= take
    by_id = {id(lot["event"]): lot for lot in lots}
    out: list[dict[str, Any]] = []
    balance = 0
    burned = 0
    expiring: list[dict[str, Any]] = []
    for e in events:
        row = dict(e)
        try:
            delta = int(e.get("delta") or 0)
        except (TypeError, ValueError):
            delta = 0
        lot = by_id.get(id(e))
        if delta > 0 and lot:
            rem = int(lot["remaining"])
            exp = int(lot["expires"] or 0)
            expired = bool(exp and exp <= now)
            live = 0 if expired else rem
            row["gross"] = delta
            row["expires"] = exp
            row["expired"] = expired
            row["remaining"] = live
            row["delta"] = live
            if expired:
                burned += rem
                row["title"] = "Сгорело"
                note = str(e.get("note") or "").strip()
                row["note"] = f"{note} · было {delta} ₽ · истекли {expire_days} дн.".strip(" ·")
            else:
                balance += rem
                if rem > 0 and exp and 0 < exp - now <= warn_days * 86400:
                    expiring.append(row)
        else:
            row["expired"] = False
            row["gross"] = delta
            row["remaining"] = delta
        out.append(row)
    return {
        "balance": max(0, balance),
        "burned": burned,
        "events": out,
        "expiring": expiring,
        "expire_days": expire_days,
        "warn_days": warn_days,
    }


def bonus_statement(contact_id: int) -> dict[str, Any]:
    key = str(int(contact_id))
    with _LOCK:
        events = list((_load()["bonuses"].get(key) or {}).get("events") or [])
    live = _live_bonus(events)
    accrued = 0
    for e in live["events"]:
        try:
            g = int(e.get("gross") or 0)
        except (TypeError, ValueError):
            g = 0
        if g > 0:
            accrued += g
    return {
        "balance": live["balance"],
        "accrued": accrued,
        "burned": live["burned"],
        "percent": int(settings.bonus_percent or 5),
        "expire_days": live["expire_days"],
        "warn_days": live["warn_days"],
        "events": live["events"],
        "expiring": live["expiring"],
    }


def take_expiry_warnings(contact_id: int) -> list[dict[str, Any]]:
    """События, по которым ещё не слали push о сгорании."""
    stmt = bonus_statement(contact_id)
    due = stmt["expiring"]
    if not due:
        return []
    with _LOCK:
        data = _load()
        warned: list[str] = data.setdefault("warned", [])
        fresh = []
        for e in due:
            wid = f"exp:{contact_id}:{e.get('id')}"
            if wid in warned:
                continue
            warned.append(wid)
            fresh.append(e)
        _save(data)
    return fresh


def touch_audience(contact_id: Any, phone: str = "", name: str = "") -> None:
    try:
        cid = int(contact_id or 0)
    except (TypeError, ValueError):
        cid = 0
    if not cid:
        return
    key = str(cid)
    with _LOCK:
        data = _load()
        row = data.setdefault("audience", {}).setdefault(key, {})
        row["contact_id"] = cid
        row["phone"] = phone or row.get("phone") or ""
        row["name"] = name or row.get("name") or ""
        row["last_seen"] = int(time.time())
        _save(data)


def list_audience() -> list[dict[str, Any]]:
    with _LOCK:
        rows = list((_load().get("audience") or {}).values())
    rows.sort(key=lambda x: int(x.get("last_seen") or 0), reverse=True)
    return rows


def upsert_identity(phone: str, *, email: str = "", contact_id: Any = None, name: str = "") -> dict[str, Any]:
    key = (phone or "").strip()
    if not key:
        return {}
    with _LOCK:
        data = _load()
        row = data.setdefault("identities", {}).setdefault(key, {})
        row["phone"] = key
        if email:
            row["email"] = email.strip().lower()
        if name:
            row["name"] = name
        try:
            cid = int(contact_id or 0)
        except (TypeError, ValueError):
            cid = 0
        if cid:
            row["contact_id"] = cid
        row["updated"] = int(time.time())
        _save(data)
        return dict(row)


def get_identity(phone: str) -> dict[str, Any] | None:
    key = (phone or "").strip()
    if not key:
        return None
    with _LOCK:
        row = (_load().get("identities") or {}).get(key)
    return dict(row) if isinstance(row, dict) else None


def identity_by_email(email: str) -> dict[str, Any] | None:
    wanted = (email or "").strip().lower()
    if not wanted:
        return None
    with _LOCK:
        rows = (_load().get("identities") or {}).values()
        for row in rows:
            if isinstance(row, dict) and str(row.get("email") or "").lower() == wanted:
                return dict(row)
    return None


def hide_car(contact_id: Any, car_id: Any) -> None:
    try:
        cid = str(int(contact_id or 0))
        kid = int(car_id)
    except (TypeError, ValueError):
        return
    if not kid:
        return
    with _LOCK:
        data = _load()
        lst = data.setdefault("hidden_cars", {}).setdefault(cid, [])
        if kid not in lst:
            lst.append(kid)
            _save(data)


def hidden_car_ids(contact_id: Any) -> set[int]:
    try:
        cid = str(int(contact_id or 0))
    except (TypeError, ValueError):
        return set()
    with _LOCK:
        rows = (_load().get("hidden_cars") or {}).get(cid) or []
    out: set[int] = set()
    for x in rows:
        try:
            out.add(int(x))
        except (TypeError, ValueError):
            continue
    return out


def save_garage_car(contact_id: Any, car: dict[str, Any]) -> None:
    try:
        cid = str(int(contact_id or 0))
        kid = int(car.get("id") or 0)
    except (TypeError, ValueError):
        return
    if not kid:
        return
    with _LOCK:
        data = _load()
        rows: list[dict[str, Any]] = data.setdefault("garage", {}).setdefault(cid, [])
        kept = [r for r in rows if int(r.get("id") or 0) != kid]
        kept.insert(0, {**car, "id": kid})
        data["garage"][cid] = kept
        _save(data)


def list_garage_local(contact_id: Any) -> list[dict[str, Any]]:
    try:
        cid = str(int(contact_id or 0))
    except (TypeError, ValueError):
        return []
    with _LOCK:
        rows = list((_load().get("garage") or {}).get(cid) or [])
    return rows


def drop_garage_car(contact_id: Any, car_id: Any) -> None:
    try:
        cid = str(int(contact_id or 0))
        kid = int(car_id)
    except (TypeError, ValueError):
        return
    with _LOCK:
        data = _load()
        rows = data.setdefault("garage", {}).setdefault(cid, [])
        data["garage"][cid] = [r for r in rows if int(r.get("id") or 0) != kid]
        _save(data)


def add_campaign(
    *,
    title: str,
    body: str,
    badge: str = "Акция",
    segment: str = "all",
    branch_id: str = "",
    pin: bool = False,
    kind: str = "push",
    image: str = "",
) -> dict[str, Any]:
    row = {
        "id": uuid.uuid4().hex[:10],
        "title": title.strip(),
        "body": body.strip(),
        "subtitle": body.strip(),
        "badge": (badge or "Акция").strip() or "Акция",
        "segment": (segment or "all").strip() or "all",
        "branch_id": str(branch_id or "").strip(),
        "pinned": bool(pin),
        "created": int(time.time()),
        "sent": 0,
        "kind": (kind or "push").strip() or "push",
        "image": (image or "").strip(),
    }
    with _LOCK:
        data = _load()
        data.setdefault("campaigns", []).insert(0, row)
        _save(data)
    return row


def mark_campaign_sent(cid: str, count: int) -> None:
    with _LOCK:
        data = _load()
        for row in data.get("campaigns") or []:
            if row.get("id") == cid:
                row["sent"] = int(count)
                row["sent_at"] = int(time.time())
                break
        _save(data)


def list_campaigns() -> list[dict[str, Any]]:
    with _LOCK:
        return list(_load().get("campaigns") or [])


def list_promos() -> list[dict[str, Any]]:
    own = [c for c in list_campaigns() if c.get("pinned") and c.get("kind") != "push"]
    if own:
        return own
    return list(DEFAULT_PROMOS)


def add_short(*, title: str, subtitle: str = "", url: str = "", poster: str = "") -> dict[str, Any]:
    row = {
        "id": uuid.uuid4().hex[:10],
        "title": title.strip(),
        "subtitle": subtitle.strip(),
        "url": url.strip(),
        "poster": poster.strip(),
        "created": int(time.time()),
    }
    with _LOCK:
        data = _load()
        data.setdefault("shorts", []).insert(0, row)
        _save(data)
    return row


def list_shorts() -> list[dict[str, Any]]:
    with _LOCK:
        data = _load()
        own = list(data.get("shorts") or [])
        hidden = {str(x) for x in (data.get("hidden_shorts") or [])}
    own_ids = {str(x.get("id")) for x in own if x.get("id")}
    defaults = [dict(d) for d in DEFAULT_SHORTS if d["id"] not in hidden and d["id"] not in own_ids]
    return own + defaults


def list_shorts_admin() -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    with _LOCK:
        data = _load()
        own = list(data.get("shorts") or [])
        hidden = {str(x) for x in (data.get("hidden_shorts") or [])}
    own_ids = {str(x.get("id")) for x in own if x.get("id")}
    for row in own:
        items.append({**row, "builtin": False})
    for d in DEFAULT_SHORTS:
        if d["id"] in own_ids or d["id"] in hidden:
            continue
        items.append({**d, "builtin": True})
    return items


def update_short(
    sid: str,
    *,
    title: str,
    subtitle: str = "",
    url: str = "",
    poster: str = "",
) -> dict[str, Any]:
    sid = (sid or "").strip()
    if not sid:
        raise KeyError("id")
    title = title.strip()
    if not title:
        raise ValueError("title")
    with _LOCK:
        data = _load()
        rows: list[dict[str, Any]] = data.setdefault("shorts", [])
        for row in rows:
            if str(row.get("id")) == sid:
                row["title"] = title
                row["subtitle"] = subtitle.strip()
                row["url"] = url.strip()
                row["poster"] = poster.strip()
                row["updated"] = int(time.time())
                _save(data)
                return dict(row)
        builtin = next((dict(d) for d in DEFAULT_SHORTS if d["id"] == sid), None)
        if not builtin:
            raise KeyError(sid)
        builtin.update(
            {
                "title": title,
                "subtitle": subtitle.strip(),
                "url": url.strip(),
                "poster": poster.strip(),
                "updated": int(time.time()),
                "created": int(time.time()),
            }
        )
        hidden = data.setdefault("hidden_shorts", [])
        if sid in hidden:
            hidden.remove(sid)
        rows.insert(0, builtin)
        _save(data)
        return builtin


def delete_short(sid: str) -> bool:
    sid = (sid or "").strip()
    if not sid:
        return False
    with _LOCK:
        data = _load()
        rows: list[dict[str, Any]] = data.setdefault("shorts", [])
        kept = [r for r in rows if str(r.get("id")) != sid]
        removed = len(kept) != len(rows)
        data["shorts"] = kept
        if any(d["id"] == sid for d in DEFAULT_SHORTS):
            hidden = data.setdefault("hidden_shorts", [])
            if sid not in hidden:
                hidden.append(sid)
            removed = True
        if not removed:
            return False
        _save(data)
    return True


def register_fcm_token(contact_id: int, token: str) -> None:
    token = (token or "").strip()
    if not token:
        return
    key = str(int(contact_id))
    with _LOCK:
        data = _load()
        book: dict[str, Any] = data.setdefault("devices", {})
        lst = [t for t in (book.get(key) or []) if t and t != token]
        lst.append(token)
        book[key] = lst[-8:]
        _save(data)


def fcm_tokens(contact_id: int) -> list[str]:
    key = str(int(contact_id))
    with _LOCK:
        return [str(t) for t in ((_load().get("devices") or {}).get(key) or []) if t]


def fcm_contact_ids() -> list[int]:
    with _LOCK:
        return [int(k) for k in (_load().get("devices") or {}) if str(k).isdigit()]
